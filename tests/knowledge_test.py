"""知识库管理模块冒烟测试：上传/解析/切片/向量化/管理/效果测试全链路。

运行方式（项目根目录）：
    .venv\\Scripts\\python.exe tests\\knowledge_test.py

前置条件：服务已启动
    .venv\\Scripts\\python.exe -m uvicorn app.main:app --port 8000
"""

from __future__ import annotations

import asyncio
import io
import os
import sys
from datetime import datetime
from pathlib import Path

import httpx

PORT = os.environ.get("CS_PORT", "8000")
BASE = f"http://127.0.0.1:{PORT}"
SAMPLES = Path(__file__).parent / "samples"
RUN_TAG = datetime.now().strftime("%m%d%H%M%S")


def make_docx_bytes() -> bytes:
    """生成带标题样式与表格的 Word 测试文档。"""
    import docx

    d = docx.Document()
    d.add_heading("售后政策", level=1)
    d.add_paragraph("电子产品整机保修一年，充电器、数据线等配件保修三个月。")
    d.add_paragraph("人为损坏不在保修范围内，维修费用以检测结果为准。")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "说明"
    table.cell(1, 0).text = "保修期"
    table.cell(1, 1).text = "整机一年"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_pdf_bytes() -> bytes:
    """生成两页 PDF：每页带重复页眉与页码，用于验证页眉页脚清理与页码元数据。"""

    def text_stream(lines: list[str]) -> bytes:
        cmds = " BT /F1 12 Tf 72 720 Td "
        for i, line in enumerate(lines):
            esc = line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
            if i == 0:
                cmds += f"({esc}) Tj "
            else:
                cmds += f"0 -20 Td ({esc}) Tj "
        cmds += " ET"
        return cmds.encode("latin-1")

    page1_body = [
        "INTERNAL ONLY",
        "Page 1",
        "Refund policy: returns accepted within 7 days after receipt.",
    ]
    page2_body = [
        "INTERNAL ONLY",
        "Page 2",
        "Shipping: most regions ship within 48 hours.",
    ]
    s1 = text_stream(page1_body)
    s2 = text_stream(page2_body)

    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R 4 0 R] /Count 2 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 5 0 R "
        b"/Resources << /Font << /F1 7 0 R >> >> >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 6 0 R "
        b"/Resources << /Font << /F1 7 0 R >> >> >>",
        b"<< /Length " + str(len(s1)).encode() + b" >>\nstream\n" + s1 + b"\nendstream",
        b"<< /Length " + str(len(s2)).encode() + b" >>\nstream\n" + s2 + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, obj in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets[1:]:
        out += f"{off:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref}\n%%EOF"
    ).encode()
    return bytes(out)


async def main() -> int:
    failures: list[str] = []

    def check(name: str, ok: bool, detail: str = "") -> None:
        mark = "[PASS]" if ok else "[FAIL]"
        print(f"  {mark} {name}" + (f"  -> {detail}" if detail else ""))
        if not ok:
            failures.append(name)

    async with httpx.AsyncClient(base_url=BASE, timeout=120) as client:
        # 1. 上传五种格式（含批量多文件）
        uploads = [
            (f"退货政策_{RUN_TAG}.md", (SAMPLES / "退货政策.md").read_bytes(), "text/markdown"),
            (f"服务说明_{RUN_TAG}.txt", (SAMPLES / "服务说明.txt").read_bytes(), "text/plain"),
            (f"帮助中心_{RUN_TAG}.html", (SAMPLES / "帮助中心.html").read_bytes(), "text/html"),
            (f"售后政策_{RUN_TAG}.docx", make_docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            (f"服务手册_{RUN_TAG}.pdf", make_pdf_bytes(), "application/pdf"),
        ]
        created: dict[str, dict] = {}
        for name, raw, mime in uploads:
            r = await client.post(
                "/api/knowledge/documents",
                files={"files": (name, raw, mime)},
            )
            data = r.json()
            item = data["results"][0]
            ok = item["ok"] and item["doc"]["status"] == "completed" and item["doc"]["chunk_count"] > 0
            check(f"上传并解析 {name}", ok, f"chunks={item.get('doc', {}).get('chunk_count')}")
            if ok:
                created[name] = item["doc"]

        # 批量上传两个 TXT
        r = await client.post(
            "/api/knowledge/documents",
            files=[
                ("files", ("批量一.txt", "批量上传测试一：客服转接规则。".encode("utf-8"), "text/plain")),
                ("files", ("批量二.txt", "批量上传测试二：会话历史保留策略。".encode("utf-8"), "text/plain")),
            ],
        )
        batch = r.json()["results"]
        check("多文件批量上传", all(b["ok"] for b in batch) and len(batch) == 2)

        # 2. 文档列表与统计
        r = await client.get("/api/knowledge/documents")
        data = r.json()
        docs = {d["filename"]: d for d in data["documents"]}
        check("文档列表与统计", data["stats"]["docs"] >= len(uploads) and data["stats"]["chunks"] > 0,
              f"docs={data['stats']['docs']}, chunks={data['stats']['chunks']}")

        # 3. PDF：页眉页脚被清理，页码元数据保留
        pdf_doc = docs.get(f"服务手册_{RUN_TAG}.pdf")
        if pdf_doc:
            r = await client.get(f"/api/knowledge/documents/{pdf_doc['id']}/chunks")
            pdf_chunks = r.json()
            no_header = all("INTERNAL ONLY" not in c["content"] and "Page " not in c["content"] for c in pdf_chunks)
            has_page2 = any(c.get("page") == 2 and "Shipping" in c["content"] for c in pdf_chunks)
            check("PDF 页眉页脚清理与页码元数据", no_header and has_page2,
                  f"chunks={len(pdf_chunks)}")

        # 4. Word：标题结构保留为切片标题
        docx_doc = docs.get(f"售后政策_{RUN_TAG}.docx")
        if docx_doc:
            r = await client.get(f"/api/knowledge/documents/{docx_doc['id']}/chunks")
            docx_chunks = r.json()
            titled = any(c.get("title") == "售后政策" and "保修" in c["content"] for c in docx_chunks)
            check("Word 标题结构保留为切片元数据", titled, f"chunks={len(docx_chunks)}")

        # 5. Markdown/HTML：标题与内容
        md_doc = docs.get(f"退货政策_{RUN_TAG}.md")
        if md_doc:
            r = await client.get(f"/api/knowledge/documents/{md_doc['id']}/chunks")
            md_chunks = r.json()
            check("Markdown 切片含标题与内容",
                  any(c.get("title") == "退货政策" and "无理由退货" in c["content"] for c in md_chunks))
        html_doc = docs.get(f"帮助中心_{RUN_TAG}.html")
        if html_doc:
            r = await client.get(f"/api/knowledge/documents/{html_doc['id']}/chunks")
            html_chunks = r.json()
            check("HTML 正文提取与标题",
                  any(c.get("title") == "账号与安全" and "密码" in c["content"] for c in html_chunks))

        # 6. 效果测试：问题命中退货政策
        r = await client.post("/api/knowledge/test", json={"question": "我想退货，钱多久能退回来？", "top_k": 5})
        data = r.json()
        results = data["results"]
        hit = any("退货政策" in res["doc_name"] for res in results)
        score_ok = all(0 < res["similarity"] <= 1 for res in results)
        check("效果测试命中相关文档", bool(results) and hit and score_ok,
              "|".join(f"{res['doc_name']}:{res['similarity']}" for res in results[:3]))

        # 7. 停用后不再命中，启用后恢复
        if md_doc:
            r = await client.patch(
                f"/api/knowledge/documents/{md_doc['id']}/enabled",
                json={"enabled": False},
            )
            check("停用文档", r.json().get("ok"))
            r = await client.post("/api/knowledge/test", json={"question": "退货政策是什么？", "top_k": 5})
            disabled_hit = any(f"退货政策_{RUN_TAG}.md" == res["doc_name"] for res in r.json()["results"])
            check("停用后不参与检索", not disabled_hit)
            r = await client.patch(
                f"/api/knowledge/documents/{md_doc['id']}/enabled",
                json={"enabled": True},
            )
            check("重新启用文档", r.json().get("ok"))

        # 8. 更新（重新解析入库）
        txt_doc = docs.get(f"服务说明_{RUN_TAG}.txt")
        if txt_doc:
            r = await client.post(f"/api/knowledge/documents/{txt_doc['id']}/reindex")
            updated = r.json()
            check("更新文档（重新解析）", updated.get("status") == "completed" and updated.get("chunk_count", 0) > 0,
                  f"chunks={updated.get('chunk_count')}")

        # 9. 删除文档
        if html_doc:
            r = await client.delete(f"/api/knowledge/documents/{html_doc['id']}")
            check("删除文档", r.status_code == 200 and r.json().get("ok"))
            r = await client.get(f"/api/knowledge/documents/{html_doc['id']}")
            check("删除后文档不存在", r.status_code == 404)

        # 10. 不支持的类型返回失败原因
        r = await client.post(
            "/api/knowledge/documents",
            files={"files": ("恶意程序.exe", b"MZ...", "application/octet-stream")},
        )
        bad = r.json()["results"][0]
        check("不支持的类型被拒绝", not bad["ok"] and "不支持" in bad["error"], bad["error"])

        # 11. 非法 top_k 校验
        r = await client.post("/api/knowledge/test", json={"question": "x", "top_k": 99})
        check("top_k 范围校验", r.status_code == 400)

    print()
    if failures:
        print(f"失败 {len(failures)} 项：{', '.join(failures)}")
        return 1
    print("全部通过 ✓")
    return 0


if __name__ == "__main__":
    if sys.platform == "win32":
        sys.stdout.reconfigure(encoding="utf-8")
    raise SystemExit(asyncio.run(main()))
